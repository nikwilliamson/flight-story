#!/usr/bin/env python3
"""Export a FlightMemory (flightmemory.com) flight history to .xlsx / .csv.

Same job as github.com/TobiasUr/FlightMemoryExporter, minus the Tk GUI:
headless Chrome logs in, walks the FLIGHTDATA list 50 legs at a time, and the rows land in a
sheet whose columns line up with Flight Log.xlsx (Date, From Code, To Code,
Airline, Flight #, Aircraft, Tail #, Logged Miles, ...).

    export FM_USERNAME=...            # or you'll be prompted
    python fm_export.py                         # -> flightmemory.xlsx
    python fm_export.py --compare "/path/to/Flight Log.xlsx"
    python fm_export.py --from-html raw/        # re-parse saved pages, no login

Credentials come from FM_USERNAME / FM_PASSWORD or an interactive prompt and
are only typed into FlightMemory's own login form. Nothing is written to disk
except the export and (with --save-html) the raw result pages.
"""
from __future__ import annotations

import argparse
import csv
import getpass
import os
import re
import sys
from dataclasses import dataclass, fields
from datetime import date, datetime
from pathlib import Path

from bs4 import BeautifulSoup, Tag

BASE_URL = "https://www.flightmemory.com"
PAGE_SIZE = 50
NEAR_DAYS = 3
KM_PER_MILE = 1.609344


@dataclass
class Flight:
    fm_no: int | None  # FlightMemory's own running leg number, 1 = first flight
    date: date | None
    date_text: str  # the date cell as shown, for partial dates that don't parse
    dep_time: str
    arr_time: str
    from_code: str
    from_name: str
    from_country: str
    to_code: str
    to_name: str
    to_country: str
    airline: str
    flight_no: str
    aircraft: str
    tail: str
    miles: int | None
    duration: str
    seat: str
    seat_type: str
    cabin: str
    role: str
    reason: str
    future: bool
    fm_id: str  # FlightMemory database id, from the row's edit link
    raw: str  # pipe-joined cell text, so nothing the parser misses is lost


# Flight Log.xlsx column names first, FlightMemory-only extras after.
COLUMNS = {
    "fm_no": "FM #",
    "date": "Date",
    "date_text": "FM Date",
    "from_name": "From",
    "from_code": "From Code",
    "to_name": "To",
    "to_code": "To Code",
    "airline": "Airline",
    "flight_no": "Flight #",
    "aircraft": "Aircraft",
    "tail": "Tail #",
    "miles": "Logged Miles",
    "from_country": "From Country",
    "to_country": "To Country",
    "dep_time": "Dep Time",
    "arr_time": "Arr Time",
    "duration": "Duration",
    "seat": "Seat",
    "seat_type": "Seat Type",
    "cabin": "Class",
    "role": "Role",
    "reason": "Reason",
    "future": "Future",
    "fm_id": "FM ID",
    "raw": "FM Raw",
}
assert set(COLUMNS) == {f.name for f in fields(Flight)}


def values(f: Flight) -> list:
    return [getattr(f, k) for k in COLUMNS]


# ---------------------------------------------------------------- scraping


def flightdata_url(dbpos: int) -> str:
    return f"{BASE_URL}/signin/?go=flugdaten&dbpos={dbpos}"


def fetch_pages(username: str, password: str, headed: bool) -> list[str]:
    from selenium import webdriver
    from selenium.common.exceptions import TimeoutException
    from selenium.webdriver.common.by import By
    from selenium.webdriver.support import expected_conditions as ec
    from selenium.webdriver.support.wait import WebDriverWait

    opts = webdriver.ChromeOptions()
    if not headed:
        opts.add_argument("--headless=new")
    driver = webdriver.Chrome(options=opts)  # Selenium Manager fetches chromedriver
    wait = WebDriverWait(driver, 15)
    try:
        driver.get(BASE_URL)
        wait.until(ec.element_to_be_clickable((By.NAME, "username"))).send_keys(username)
        wait.until(ec.element_to_be_clickable((By.NAME, "passwort"))).send_keys(password)
        wait.until(ec.element_to_be_clickable(
            (By.XPATH, "//input[@type='submit' and @value='SignIn']"))).click()
        try:
            wait.until(ec.presence_of_element_located(
                (By.XPATH, "//a[normalize-space()='FLIGHTDATA']")))
        except TimeoutException:
            sys.exit("Login failed: FLIGHTDATA link never appeared (check username/password).")

        # The list pages 50 legs at a time via ?dbpos=; walk it directly rather
        # than clicking the pager, until a page comes back empty or repeats.
        pages: list[str] = []
        seen: set[int | None] = set()
        for dbpos in range(0, 100_000, PAGE_SIZE):
            driver.get(flightdata_url(dbpos))
            try:
                wait.until(lambda d: count_rows(d.page_source) > 0)
            except TimeoutException:
                break
            html = driver.page_source
            numbers = {row_number(tr) for tr in flight_rows(html)}
            if numbers <= seen:
                break
            seen |= numbers
            pages.append(html)
            print(f"page {len(pages)}: {len(numbers)} flights", file=sys.stderr)
        if not pages:
            print("warning: FLIGHTDATA had no flight rows", file=sys.stderr)
            pages.append(driver.page_source)
        return pages
    finally:
        driver.quit()


# ---------------------------------------------------------------- parsing


DATE = re.compile(r"(\d{1,2})[-./](\d{1,2})[-./](\d{4})")
TIME = re.compile(r"\b(\d{1,2}:\d{2})\b")
NUMBER = re.compile(r"[\d.,]+")
EDIT_ID = re.compile(r"[?&]id=(\d+)")


def cells(tr: Tag) -> list[Tag]:
    # The distance/duration cell is a <th>, so count both.
    return tr.find_all(["td", "th"], recursive=False)


def flight_rows(html: str) -> list[Tag]:
    """Flight list rows: 10+ cells, the first being the leg number.

    Not keyed on the date: older legs can carry partial dates (month or year only).
    """
    soup = BeautifulSoup(html, "html.parser")
    root = soup.select_one(".container") or soup
    return [tr for tr in root.find_all("tr")
            if len(c := cells(tr)) >= 10 and c[0].get_text(strip=True).isdigit()]


def count_rows(html: str) -> int:
    return len(flight_rows(html))


def lines(tag: Tag) -> list[str]:
    """Text split on <br>/element boundaries."""
    return [s for s in tag.get_text("|", strip=True).split("|") if s]


def row_number(tr: Tag) -> int | None:
    text = cells(tr)[0].get_text(strip=True)
    return int(text) if text.isdigit() else None


def to_int(text: str) -> int | None:
    m = NUMBER.search(text)
    return int(re.sub(r"[.,]", "", m.group())) if m else None


def parse_distance(th: Tag) -> tuple[int | None, str]:
    """'940|mi|2:11|h' -> (940, '2:11'); km converted to miles."""
    parts = lines(th)
    miles, duration = None, ""
    for value, unit in zip(parts, parts[1:]):
        if unit in ("mi", "km") and (n := to_int(value)) is not None:
            miles = n if unit == "mi" else round(n / KM_PER_MILE)
        elif unit == "h" and TIME.fullmatch(value):
            duration = value
    return miles, duration


def parse_seat(td: Tag) -> tuple[str, str, str, str, str]:
    """'12A / Window' + <small>Economy|Passenger|Personal</small> -> its five parts.

    The <small> block is [class], role, reason; class is omitted when unset.
    """
    small = td.find("small")
    details = lines(small) if small else []
    if small:
        small.extract()
    seat, _, seat_type = td.get_text(" ", strip=True).partition("/")
    cabin = details[0] if len(details) >= 3 else ""
    role, reason = details[-2:] if len(details) >= 2 else ("", "")
    cabin = {"EconomyPlus": "Economy Plus"}.get(cabin, cabin)
    return seat.strip(), seat_type.strip(), cabin, role, reason


def parse_row(tr: Tag, day_first: bool) -> Flight:
    c = cells(tr)
    when = c[1].get_text(" ", strip=True)
    d = DATE.search(when)
    times = TIME.findall(when)
    from_city, from_country, *_ = lines(c[3]) + ["", ""]
    to_city, to_country, *_ = lines(c[5]) + ["", ""]
    airline, flight_no, *_ = lines(c[7]) + ["", ""]
    aircraft, tail, *_ = lines(c[8]) + ["", ""]
    miles, duration = parse_distance(c[6])
    raw = " | ".join(x.get_text(" ", strip=True) for x in c[:10])
    seat, seat_type, cabin, role, reason = parse_seat(c[9])
    edit = tr.find("option", value=EDIT_ID)
    return Flight(
        fm_no=row_number(tr),
        date=parse_date(d, day_first),
        date_text=when,
        dep_time=times[0] if times else "",
        arr_time=times[1] if len(times) > 1 else "",
        from_code=c[2].get_text(strip=True),
        from_name=from_city,
        from_country=from_country,
        to_code=c[4].get_text(strip=True),
        to_name=to_city,
        to_country=to_country,
        airline=airline,
        flight_no=flight_no,
        aircraft=aircraft,
        tail=tail,
        miles=miles,
        duration=duration,
        seat=seat,
        seat_type=seat_type,
        cabin=cabin,
        role=role,
        reason=reason,
        future=tr.get("title") == "Future Flight",
        fm_id=EDIT_ID.search(edit["value"])[1] if edit else "",
        raw=raw,
    )


def parse_date(m: re.Match | None, day_first: bool) -> date | None:
    if not m:
        return None
    a, b, year = int(m[1]), int(m[2]), int(m[3])
    day, month = (a, b) if day_first else (b, a)
    try:
        return date(year, month, day)
    except ValueError:  # e.g. 00-00-1985 for an unknown day
        return None


def is_day_first(rows: list[Tag]) -> bool:
    """FlightMemory renders dates in the account's locale (mm-dd-yyyy or dd.mm.yyyy)."""
    pairs = [(int(m[1]), int(m[2])) for tr in rows
             if (m := DATE.search(cells(tr)[1].get_text(" ", strip=True)))]
    if any(a > 12 for a, _ in pairs):
        return True
    if any(b > 12 for _, b in pairs):
        return False
    return any("." in cells(tr)[1].get_text() for tr in rows)


def parse_pages(pages: list[str]) -> list[Flight]:
    rows = [tr for html in pages for tr in flight_rows(html)]
    day_first = is_day_first(rows)
    by_number: dict[int | None, Flight] = {}
    unnumbered: list[Flight] = []
    for tr in rows:
        f = parse_row(tr, day_first)
        if f.fm_no is None:
            unnumbered.append(f)
        else:
            by_number[f.fm_no] = f  # pages can overlap; last copy wins
    # FM # is FlightMemory's chronological leg number, the same order as the log.
    return sorted(by_number.values(), key=lambda f: f.fm_no) + unnumbered


# ---------------------------------------------------------------- output


def write_xlsx(path: Path, flights: list[Flight], reconcile: dict[str, list[list]] | None) -> None:
    import openpyxl
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FlightMemory"
    ws.append(list(COLUMNS.values()))
    for f in flights:
        ws.append(values(f))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for name, rows in (reconcile or {}).items():
        sheet = wb.create_sheet(name)
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        if sheet.max_row > 1:
            sheet.auto_filter.ref = sheet.dimensions

    for sheet in wb.worksheets:
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, date):
                    cell.number_format = "yyyy-mm-dd"
        for i, col in enumerate(sheet.iter_cols(max_row=min(sheet.max_row, 200)), start=1):
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            sheet.column_dimensions[get_column_letter(i)].width = min(width + 2, 40)
    wb.save(path)


def write_csv(path: Path, flights: list[Flight]) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(COLUMNS.values())
        for f in flights:
            w.writerow(values(f))


# ---------------------------------------------------------------- reconcile


def as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    return value if isinstance(value, date) else None


def norm(value) -> str:
    return str(value or "").replace(" ", "").upper()


def reconcile(flights: list[Flight], log_path: Path) -> dict[str, list[list]]:
    """Pair FlightMemory legs with Flight Log rows and list what differs.

    Pass 1 pairs on exact (date, from code, to code). Pass 2 pairs what's left
    when the flight number matches on the same date (an airport code differs), or
    the route matches within NEAR_DAYS (the date differs), nearest date first.
    """
    import openpyxl

    ws = openpyxl.load_workbook(log_path, read_only=True, data_only=True)["Flights"]
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))
    log = [dict(zip(header, r)) for r in rows if r[header.index("ID")] is not None]
    for r in log:
        r["Date"] = as_date(r["Date"])

    pairs: dict[int, dict] = {}  # FM list index -> log row
    taken: set[int] = set()

    def claim(i: int, r: dict) -> None:
        pairs[i] = r
        taken.add(r["ID"])

    exact: dict[tuple, list[dict]] = {}
    for r in log:
        exact.setdefault((r["Date"], norm(r["From Code"]), norm(r["To Code"])), []).append(r)
    for i, f in enumerate(flights):
        free = [r for r in exact.get((f.date, norm(f.from_code), norm(f.to_code)), [])
                if r["ID"] not in taken]
        if free:
            claim(i, free[0])

    for i, f in enumerate(flights):
        if i in pairs or not f.date:
            continue
        candidates = [
            r for r in log
            if r["ID"] not in taken and r["Date"] and (
                (r["Date"] == f.date and f.flight_no and norm(r["Flight #"]) == norm(f.flight_no))
                or (norm(r["From Code"]) == norm(f.from_code) and norm(r["To Code"]) == norm(f.to_code)
                    and abs((r["Date"] - f.date).days) <= NEAR_DAYS))]
        if candidates:
            claim(i, min(candidates, key=lambda r: abs((r["Date"] - f.date).days)))

    matched = [["Log ID", "FM #", "Date", "From", "To", "Field", "Flight Log", "FlightMemory"]]
    fm_only = [list(COLUMNS.values())]
    for i, f in enumerate(flights):
        r = pairs.get(i)
        if r is None:
            fm_only.append(values(f))
            continue
        base = [r["ID"], f.fm_no, f.date, f.from_code, f.to_code]
        diffs = [(label, r[label], fm) for label, fm in
                 (("Date", f.date), ("From Code", f.from_code), ("To Code", f.to_code),
                  ("Flight #", f.flight_no), ("Tail #", f.tail))
                 if fm and (r[label] != fm if label == "Date" else norm(r[label]) != norm(fm))]
        matched += [base + list(d) for d in diffs] or [base + ["(all match)", "", ""]]

    dated = [f.date for f in flights if f.date]
    lo, hi = (min(dated), max(dated)) if dated else (None, None)
    log_cols = ["ID", "Date", "From Code", "To Code", "Airline", "Flight #", "FM Sync"]
    log_only = [["Log ID"] + log_cols[1:]]
    log_only += [[r[c] for c in log_cols] for r in log
                 if r["ID"] not in taken and lo and r["Date"] and lo <= r["Date"] <= hi]

    print(f"reconcile: {len(pairs)} paired, {len(fm_only) - 1} only on FlightMemory, "
          f"{len(log_only) - 1} only in the log ({lo} to {hi})", file=sys.stderr)
    return {"Matched": matched, "Only on FlightMemory": fm_only, "Only in Log": log_only}


# ---------------------------------------------------------------- main


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("-o", "--out", type=Path, default=Path("flightmemory.xlsx"),
                    help="output .xlsx or .csv (default: flightmemory.xlsx)")
    ap.add_argument("--compare", type=Path, metavar="FLIGHT_LOG_XLSX",
                    help="add Matched / Only on FlightMemory / Only in Log sheets")
    ap.add_argument("--save-html", type=Path, metavar="DIR", help="keep the raw result pages")
    ap.add_argument("--from-html", type=Path, metavar="DIR", help="parse saved pages, skip login")
    ap.add_argument("--headed", action="store_true", help="show the browser window")
    args = ap.parse_args()

    if args.from_html:
        pages = [p.read_text(encoding="utf-8") for p in sorted(args.from_html.glob("page-*.html"))]
        if not pages:
            sys.exit(f"No page-*.html files in {args.from_html}")
    else:
        username = os.environ.get("FM_USERNAME") or input("FlightMemory username: ")
        password = os.environ.get("FM_PASSWORD") or getpass.getpass("FlightMemory password: ")
        pages = fetch_pages(username, password, args.headed)

    if args.save_html:
        args.save_html.mkdir(parents=True, exist_ok=True)
        for i, html in enumerate(pages, start=1):
            (args.save_html / f"page-{i:03d}.html").write_text(html, encoding="utf-8")

    flights = parse_pages(pages)
    if not flights:
        sys.exit("Parsed 0 flights. Re-run with --save-html raw/ and check the page layout.")

    if args.out.suffix.lower() == ".csv":
        if args.compare:
            sys.exit("--compare needs .xlsx output")
        write_csv(args.out, flights)
    else:
        write_xlsx(args.out, flights, reconcile(flights, args.compare) if args.compare else None)
    print(f"{len(flights)} flights -> {args.out}", file=sys.stderr)


if __name__ == "__main__":
    main()
